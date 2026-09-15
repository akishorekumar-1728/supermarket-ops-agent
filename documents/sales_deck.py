"""
documents/sales_deck.py
=======================
PPTX Sales Analysis Deck Generator using python-pptx, pandas, and matplotlib.

Generates a PowerPoint deck (.pptx) summarizing sales, GST, payment methods,
top products, stock health, and automated business insights.
"""
from __future__ import annotations

import datetime
import io
import os
import sqlite3
from pathlib import Path
from typing import Any

import matplotlib
matplotlib.use("Agg")  # Non-interactive backend
import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from tools.inventory import get_low_stock
from tools.preferences import get_preference

GENERATED_DIR = Path("generated").resolve()

# Color Palette Constants
DARK_BLUE = RGBColor(26, 54, 93)     # #1A365D
MED_BLUE = RGBColor(43, 108, 176)    # #2B6CB0
LIGHT_BG = RGBColor(235, 248, 255)   # #EBF8FF
TEXT_DARK = RGBColor(45, 55, 72)     # #2D3748
TEXT_MUTED = RGBColor(113, 128, 150) # #718096


def _create_pie_chart_image(df: pd.DataFrame, label_col: str, val_col: str, title: str) -> io.BytesIO:
    """Generate a pie chart using matplotlib and return as in-memory BytesIO image."""
    fig, ax = plt.subplots(figsize=(6, 4.5), dpi=150)
    colors_list = ['#2B6CB0', '#4299E1', '#63B3ED', '#90CDF4', '#EBF8FF']
    
    wedges, texts, autotexts = ax.pie(
        df[val_col],
        labels=df[label_col],
        autopct='%1.1f%%',
        startangle=140,
        colors=colors_list[:len(df)],
        textprops=dict(color='#2D3748', fontsize=10),
    )
    for autotext in autotexts:
        autotext.set_color('white')
        autotext.set_weight('bold')
        
    ax.set_title(title, fontsize=12, fontweight='bold', color='#1A365D', pad=15)
    plt.tight_layout()
    
    img_buf = io.BytesIO()
    plt.savefig(img_buf, format='png', bbox_inches='tight')
    plt.close(fig)
    img_buf.seek(0)
    return img_buf


def _create_bar_chart_image(df: pd.DataFrame, label_col: str, val_col: str, title: str, xlabel: str, ylabel: str) -> io.BytesIO:
    """Generate a horizontal bar chart using matplotlib and return as in-memory BytesIO image."""
    fig, ax = plt.subplots(figsize=(6.5, 4.5), dpi=150)
    
    bars = ax.barh(df[label_col], df[val_col], color='#2B6CB0', height=0.6)
    ax.invert_yaxis()  # top-down ranking
    
    ax.set_title(title, fontsize=12, fontweight='bold', color='#1A365D', pad=15)
    ax.set_xlabel(xlabel, fontsize=10, color='#4A5568')
    ax.set_ylabel(ylabel, fontsize=10, color='#4A5568')
    
    # Add value labels to bars
    for bar in bars:
        width = bar.get_width()
        ax.text(
            width + (max(df[val_col]) * 0.02),
            bar.get_y() + bar.get_height() / 2,
            f'{width:g}',
            ha='left',
            va='center',
            fontsize=9,
            color='#2D3748',
            fontweight='bold'
        )

    # Style grid
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.grid(axis='x', linestyle='--', alpha=0.5)
    
    plt.tight_layout()
    
    img_buf = io.BytesIO()
    plt.savefig(img_buf, format='png', bbox_inches='tight')
    plt.close(fig)
    img_buf.seek(0)
    return img_buf


def generate_sales_deck(
    period: str = "week",
    conn: sqlite3.Connection | None = None,
    output_dir: str | Path | None = None,
    **kwargs: Any,
) -> str:
    """
    Generate a PowerPoint sales analysis presentation (.pptx) from SQLite data.

    Parameters:
        period: Time window string ('day', 'week', 'month', or 'all'). Default 'week'.
        conn: Open sqlite3.Connection.
        output_dir: Output directory for saving .pptx. Defaults to 'generated/'.

    Returns:
        str: Absolute file path to the created presentation file.
    """
    c = conn or kwargs.get("conn")
    if c is None:
        raise ValueError("generate_sales_deck requires an open sqlite3.Connection (conn).")

    # Date filter calculation
    now = datetime.datetime.now(datetime.timezone.utc)
    if period == "day":
        start_date = now.strftime("%Y-%m-%d")
    elif period == "month":
        start_date = (now - datetime.timedelta(days=30)).strftime("%Y-%m-%d")
    elif period == "all":
        start_date = "1970-01-01"
    else:  # 'week' default
        start_date = (now - datetime.timedelta(days=7)).strftime("%Y-%m-%d")

    # Fetch data using pandas
    bills_df = pd.read_sql_query(
        """
        SELECT id, invoice_number, payment_mode, subtotal, cgst_total, sgst_total,
               gst_total, grand_total, COALESCE(finalized_at, created_at) AS date_time
        FROM bills
        WHERE status = 'finalized'
          AND date(COALESCE(finalized_at, created_at)) >= date(?)
        """,
        c,
        params=(start_date,),
    )

    items_df = pd.read_sql_query(
        """
        SELECT bi.bill_id, bi.product_id, bi.quantity, bi.total,
               p.name AS product_name, p.unit AS unit, p.sku, p.cost_price, p.selling_price
        FROM bill_items bi
        JOIN bills b ON b.id = bi.bill_id
        JOIN products p ON p.id = bi.product_id
        WHERE b.status = 'finalized'
          AND date(COALESCE(b.finalized_at, b.created_at)) >= date(?)
        """,
        c,
        params=(start_date,),
    )

    low_stock_list = get_low_stock(c)
    store_name = get_preference("store_name", default="Supermarket Kirana Store", conn=c)

    # Calculate metrics
    total_sales = float(bills_df["grand_total"].sum()) if not bills_df.empty else 0.0
    total_bills = len(bills_df)
    avg_bill_val = (total_sales / total_bills) if total_bills > 0 else 0.0
    
    total_gst = float(bills_df["gst_total"].sum()) if not bills_df.empty else 0.0
    total_cgst = float(bills_df["cgst_total"].sum()) if not bills_df.empty else 0.0
    total_sgst = float(bills_df["sgst_total"].sum()) if not bills_df.empty else 0.0

    # Build PPTX Presentation
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # Blank layout for custom design

    def add_header(slide, title_text: str, category_text: str = "PERFORMANCE ANALYTICS"):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(1.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = MED_BLUE
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = DARK_BLUE

    # -------------------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    tx_title = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(8.0), Inches(2.5))
    tf1 = tx_title.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = f"{store_name}"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE
    
    p_sub = tf1.add_paragraph()
    p_sub.text = f"Sales Analysis & Operations Report ({period.capitalize()}ly)"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = MED_BLUE
    p_sub.space_before = Pt(10)

    p_date = tf1.add_paragraph()
    p_date.text = f"Generated: {now.strftime('%d %B %Y')} | Data Since: {start_date}"
    p_date.font.size = Pt(12)
    p_date.font.color.rgb = TEXT_MUTED
    p_date.space_before = Pt(20)

    # -------------------------------------------------------------------------
    # SLIDE 2: Sales Summary Slide
    # -------------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Sales Performance Summary")

    cards = [
        ("TOTAL SALES", f"₹ {total_sales:,.2f}"),
        ("TOTAL BILLS", f"{total_bills}"),
        ("AVERAGE BILL VALUE", f"₹ {avg_bill_val:,.2f}"),
    ]
    for idx, (label, val_str) in enumerate(cards):
        left = Inches(0.8 + idx * 2.9)
        top = Inches(2.2)
        width = Inches(2.6)
        height = Inches(2.2)

        # Draw card box
        shape = slide2.shapes.add_shape(1, left, top, width, height)  # 1 = MSO_SHAPE.RECTANGLE
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_BG
        shape.line.color.rgb = MED_BLUE

        tf_card = shape.text_frame
        tf_card.word_wrap = True
        
        p_lbl = tf_card.paragraphs[0]
        p_lbl.text = label
        p_lbl.font.size = Pt(11)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = MED_BLUE
        p_lbl.alignment = PP_ALIGN.CENTER

        p_val = tf_card.add_paragraph()
        p_val.text = val_str
        p_val.font.size = Pt(22)
        p_val.font.bold = True
        p_val.font.color.rgb = DARK_BLUE
        p_val.alignment = PP_ALIGN.CENTER
        p_val.space_before = Pt(25)

    # -------------------------------------------------------------------------
    # SLIDE 3: GST Summary Slide
    # -------------------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "GST Tax Collections Breakdown")

    gst_cards = [
        ("TOTAL GST COLLECTED", f"₹ {total_gst:,.2f}"),
        ("CGST TOTAL (CENTRAL)", f"₹ {total_cgst:,.2f}"),
        ("SGST TOTAL (STATE)", f"₹ {total_sgst:,.2f}"),
    ]
    for idx, (label, val_str) in enumerate(gst_cards):
        left = Inches(0.8 + idx * 2.9)
        top = Inches(2.2)
        width = Inches(2.6)
        height = Inches(2.2)

        shape = slide3.shapes.add_shape(1, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_BG
        shape.line.color.rgb = MED_BLUE

        tf_card = shape.text_frame
        tf_card.word_wrap = True

        p_lbl = tf_card.paragraphs[0]
        p_lbl.text = label
        p_lbl.font.size = Pt(11)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = MED_BLUE
        p_lbl.alignment = PP_ALIGN.CENTER

        p_val = tf_card.add_paragraph()
        p_val.text = val_str
        p_val.font.size = Pt(22)
        p_val.font.bold = True
        p_val.font.color.rgb = DARK_BLUE
        p_val.alignment = PP_ALIGN.CENTER
        p_val.space_before = Pt(25)

    # -------------------------------------------------------------------------
    # SLIDE 4: Payment Breakdown Slide (with Matplotlib Chart)
    # -------------------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Payment Mode Breakdown")

    if not bills_df.empty:
        pay_df = bills_df.groupby("payment_mode")["grand_total"].agg(["sum", "count"]).reset_index()
        pay_df.columns = ["Payment Mode", "Total Sales", "Bill Count"]
        pay_df["Payment Mode"] = pay_df["Payment Mode"].str.upper()
        
        # Embed Pie Chart
        chart_buf = _create_pie_chart_image(pay_df, "Payment Mode", "Total Sales", "Sales Distribution by Payment Method")
        slide4.shapes.add_picture(chart_buf, Inches(0.8), Inches(1.8), width=Inches(4.5))

        # Right side text table/summary
        tx_box = slide4.shapes.add_textbox(Inches(5.6), Inches(1.8), Inches(3.8), Inches(4.5))
        tf_pay = tx_box.text_frame
        tf_pay.word_wrap = True
        
        p0 = tf_pay.paragraphs[0]
        p0.text = "Summary Table"
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = DARK_BLUE

        for _, row in pay_df.iterrows():
            p_row = tf_pay.add_paragraph()
            p_row.text = f"• {row['Payment Mode']}: ₹ {row['Total Sales']:,.2f} ({int(row['Bill Count'])} bills)"
            p_row.font.size = Pt(12)
            p_row.font.color.rgb = TEXT_DARK
            p_row.space_before = Pt(8)
    else:
        tx_box = slide4.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(8.0), Inches(2.0))
        tx_box.text_frame.text = "No payment transaction data available for this period."

    # -------------------------------------------------------------------------
    # SLIDE 5: Top Selling Products (with Matplotlib Bar Chart)
    # -------------------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Top Selling Products")

    if not items_df.empty:
        prod_df = items_df.groupby(["product_name", "unit"])["quantity"].sum().reset_index()
        prod_df = prod_df.sort_values(by="quantity", ascending=False).head(5)

        chart_buf2 = _create_bar_chart_image(
            prod_df, "product_name", "quantity",
            "Top 5 Products by Quantity Sold", "Units Sold", "Product"
        )
        slide5.shapes.add_picture(chart_buf2, Inches(1.5), Inches(1.8), width=Inches(6.5))
    else:
        tx_box = slide5.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(8.0), Inches(2.0))
        tx_box.text_frame.text = "No product sales data available for this period."

    # -------------------------------------------------------------------------
    # SLIDE 6: Stock Health Slide
    # -------------------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Stock Health & Inventory Alerts")

    tx_stock = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.5), Inches(4.5))
    tf_stock = tx_stock.text_frame
    tf_stock.word_wrap = True

    if low_stock_list:
        p_alert = tf_stock.paragraphs[0]
        p_alert.text = f"⚠️ Low Stock Alert ({len(low_stock_list)} products require reordering):"
        p_alert.font.size = Pt(13)
        p_alert.font.bold = True
        p_alert.font.color.rgb = RGBColor(197, 48, 48)  # Red alert

        for item in low_stock_list[:6]:
            p_item = tf_stock.add_paragraph()
            p_item.text = f"• {item['name']} (SKU: {item['sku']}): Stock = {item['quantity']} {item['unit']} (Reorder Level = {item['reorder_level']}, Shortfall = {item['shortfall']})"
            p_item.font.size = Pt(11)
            p_item.font.color.rgb = TEXT_DARK
            p_item.space_before = Pt(6)
    else:
        p_ok = tf_stock.paragraphs[0]
        p_ok.text = "✅ Stock Health Good: All inventory items are above their reorder levels."
        p_ok.font.size = Pt(14)
        p_ok.font.bold = True
        p_ok.font.color.rgb = RGBColor(56, 161, 105)

    # -------------------------------------------------------------------------
    # SLIDE 7: Automated Business Insights Slide
    # -------------------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Automated Business Insights")

    insights = []
    if not items_df.empty:
        top_prod = items_df.groupby("product_name")["quantity"].sum().idxmax()
        insights.append(f"Highest demand product: '{top_prod}' leads total volume sales.")

    if not bills_df.empty:
        busiest_mode = bills_df["payment_mode"].value_counts().idxmax().upper()
        insights.append(f"Preferred payment method: '{busiest_mode}' is the most frequently used payment channel.")

    if low_stock_list:
        insights.append(f"Reorder priority: {len(low_stock_list)} products have hit inventory reorder thresholds.")
    else:
        insights.append("Inventory status: Inventory levels are well-balanced across all categories.")

    insights.append(f"Revenue per ticket: Average customer transaction value is ₹ {avg_bill_val:,.2f}.")

    tx_ins = slide7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.5), Inches(4.5))
    tf_ins = tx_ins.text_frame
    tf_ins.word_wrap = True

    p_head = tf_ins.paragraphs[0]
    p_head.text = "Key Operational Takeaways:"
    p_head.font.size = Pt(14)
    p_head.font.bold = True
    p_head.font.color.rgb = DARK_BLUE

    for ins in insights:
        p_bullet = tf_ins.add_paragraph()
        p_bullet.text = f"• {ins}"
        p_bullet.font.size = Pt(12)
        p_bullet.font.color.rgb = TEXT_DARK
        p_bullet.space_before = Pt(10)

    # Save presentation
    out_dir = Path(output_dir) if output_dir else GENERATED_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    
    file_path = out_dir / f"sales_deck_{period}_{now.strftime('%Y%m%d')}.pptx"
    prs.save(str(file_path))

    return str(file_path.resolve())
